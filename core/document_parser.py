"""
core/document_parser.py — Unified extractor for PDF, Word (.docx), and plain text documents.
"""

from __future__ import annotations

import os
import logging
from typing import Any

logger = logging.getLogger(__name__)


def parse_pdf(file_path: str, max_pages: int = 30) -> dict[str, Any]:
    """
    Extract text and metadata from a PDF file using pypdf.
    """
    try:
        import pypdf
    except ImportError:
        return {"ok": False, "error": "Missing dependency: pypdf is not installed."}

    if not os.path.exists(file_path):
        return {"ok": False, "error": f"File not found: {file_path}"}

    try:
        reader = pypdf.PdfReader(file_path)
        total_pages = len(reader.pages)
        pages_to_read = min(total_pages, max_pages)

        extracted_pages = []
        total_chars = 0

        for i in range(pages_to_read):
            page = reader.pages[i]
            page_text = (page.extract_text() or "").strip()
            total_chars += len(page_text)
            extracted_pages.append(f"### [第 {i + 1} 页 / 共 {total_pages} 页]\n{page_text}")

        content = "\n\n".join(extracted_pages)
        if pages_to_read < total_pages:
            content += f"\n\n*(文档较长，已截取前 {pages_to_read} 页，剩余 {total_pages - pages_to_read} 页未展示)*"

        is_scanned = total_chars < (pages_to_read * 30)

        meta = reader.metadata or {}
        title = meta.title or os.path.basename(file_path)

        return {
            "ok": True,
            "type": "pdf",
            "title": str(title),
            "page_count": total_pages,
            "read_pages": pages_to_read,
            "total_chars": total_chars,
            "is_scanned": is_scanned,
            "content": content,
        }
    except Exception as e:
        logger.exception("Failed to parse PDF %s: %s", file_path, e)
        return {"ok": False, "error": f"PDF 解析失败: {str(e)}"}


def parse_docx(file_path: str) -> dict[str, Any]:
    """
    Extract headings, paragraphs, and tables from a Word (.docx) file using python-docx.
    """
    try:
        import docx
    except ImportError:
        return {"ok": False, "error": "Missing dependency: python-docx is not installed."}

    if not os.path.exists(file_path):
        return {"ok": False, "error": f"File not found: {file_path}"}

    try:
        doc = docx.Document(file_path)
        chunks = []
        para_count = 0
        table_count = 0

        # Extract body elements (paragraphs and tables)
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            para_count += 1
            style_name = p.style.name.lower() if p.style and p.style.name else ""

            if "heading 1" in style_name or "title" in style_name:
                chunks.append(f"# {text}")
            elif "heading 2" in style_name:
                chunks.append(f"## {text}")
            elif "heading 3" in style_name:
                chunks.append(f"### {text}")
            elif "list" in style_name:
                chunks.append(f"- {text}")
            else:
                chunks.append(text)

        for table in doc.tables:
            table_count += 1
            rows_data = []
            for row in table.rows:
                row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                rows_data.append(row_cells)

            if rows_data:
                header = rows_data[0]
                md_table = [f"| {' | '.join(header)} |"]
                md_table.append(f"| {' | '.join(['---'] * len(header))} |")
                for row in rows_data[1:]:
                    if len(row) < len(header):
                        row.extend([""] * (len(header) - len(row)))
                    md_table.append(f"| {' | '.join(row[:len(header)])} |")
                chunks.append("\n" + "\n".join(md_table) + "\n")

        content = "\n\n".join(chunks)
        title = os.path.basename(file_path)

        return {
            "ok": True,
            "type": "docx",
            "title": title,
            "paragraphs_count": para_count,
            "tables_count": table_count,
            "content": content or "（文档内容为空）",
        }
    except Exception as e:
        logger.exception("Failed to parse docx %s: %s", file_path, e)
        return {"ok": False, "error": f"Word 文档解析失败: {str(e)}"}


def parse_text_file(file_path: str, max_chars: int = 50000) -> dict[str, Any]:
    """
    Extract content from plain text, markdown, CSV, or code files.
    """
    if not os.path.exists(file_path):
        return {"ok": False, "error": f"File not found: {file_path}"}

    encodings = ["utf-8", "utf-8-sig", "gb18030", "gbk", "latin-1"]
    content = ""
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc) as f:
                content = f.read(max_chars + 1000)
            break
        except UnicodeDecodeError:
            continue
        except Exception as e:
            return {"ok": False, "error": str(e)}

    if not content:
        return {"ok": False, "error": "无法使用常见编码读取此文本文件。"}

    is_truncated = len(content) > max_chars
    if is_truncated:
        content = content[:max_chars] + f"\n\n*(文件过大，已截取前 {max_chars} 字符)*"

    return {
        "ok": True,
        "type": "text",
        "title": os.path.basename(file_path),
        "content": content,
    }


def parse_document(file_path: str, max_pages: int = 30) -> dict[str, Any]:
    """
    Smart document parser: Automatically dispatches based on file extension.
    """
    if not os.path.exists(file_path):
        return {"ok": False, "error": f"File not found: {file_path}"}

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return parse_pdf(file_path, max_pages=max_pages)
    elif ext in (".docx", ".doc"):
        if ext == ".doc":
            return {
                "ok": False,
                "error": "抱歉，旧版二进制 .doc 格式不支持直接解析，请将文件保存为现代 .docx 格式后重试。",
            }
        return parse_docx(file_path)
    elif ext in (".txt", ".md", ".json", ".csv", ".log", ".py", ".html", ".xml", ".yaml", ".yml"):
        return parse_text_file(file_path)
    else:
        # Try text fallback
        return parse_text_file(file_path)
