import unittest
from unittest.mock import patch, MagicMock
import os
import tempfile

from core.recording_message import RecordingMessage
from core.message import Message
from core.message_request import MessageRequest
from plugins.document_reader import bot_execute
from agent.superuser_tools import read_document_executor


class TestDocumentReaderSecurity(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.temp_db_file = tempfile.NamedTemporaryFile(suffix=".sqlite", delete=False)
        self.temp_db_file.close()
        
        from store.database import Database
        from store.file_store import FileStore
        from runtime import context
        
        self.db = Database(self.temp_db_file.name)
        self.file_store = FileStore(self.db, download_dir=self.temp_dir)
        context.file_store = self.file_store

    def tearDown(self):
        try:
            os.remove(self.temp_db_file.name)
        except Exception:
            pass

    def test_message_request_files_field(self):
        req = MessageRequest({"args": "doc", "files": [{"file_name": "test.pdf"}]})
        self.assertEqual(len(req.files), 1)
        self.assertEqual(req.files[0]["file_name"], "test.pdf")

        msg = RecordingMessage({
            "frontend": "telegram",
            "context": {"group_id": "1", "user_id": "1"},
            "request": {"command": "doc", "args": "", "files": [{"file_name": "test.pdf"}]}
        })
        self.assertEqual(len(msg.request.files), 1)
        d = msg.to_dict()
        self.assertIn("files", d["request"])
        self.assertEqual(len(d["request"]["files"]), 1)

    @patch("plugins.document_reader.is_superuser")
    def test_plugin_non_superuser_denied(self, mock_su):
        mock_su.return_value = False

        msg = RecordingMessage({
            "frontend": "telegram",
            "context": {"group_id": "1", "user_id": "normal_user"},
            "request": {"command": "doc", "args": "report.pdf", "files": []}
        })
        bot_execute(msg, {})

        self.assertEqual(len(msg.outbox), 1)
        self.assertIn("403: nemo: 由于安全风险，非管理员无法处理未知文件", msg.outbox[0].text)

    @patch("agent.superuser_tools.is_superuser")
    def test_agent_tool_non_superuser_denied(self, mock_su):
        mock_su.return_value = False

        msg = Message({
            "frontend": "telegram",
            "context": {"group_id": "1", "user_id": "normal_user"},
            "request": {"command": "read_document", "args": "", "files": []}
        })
        res = read_document_executor({"query": "report.pdf"}, msg)
        self.assertIn("error", res)
        self.assertIn("非管理员无法处理未知文件", res["error"])

    @patch("plugins.document_reader.is_superuser")
    def test_plugin_superuser_allowed(self, mock_su):
        mock_su.return_value = True

        # Create dummy docx
        import docx
        doc_path = os.path.join(self.temp_dir, "meeting.docx")
        d = docx.Document()
        d.add_heading("会议纪要", level=1)
        d.add_paragraph("讨论重点：提升系统鲁棒性。")
        d.save(doc_path)

        msg = RecordingMessage({
            "frontend": "telegram",
            "context": {"group_id": "1", "user_id": "super_admin"},
            "request": {
                "command": "doc",
                "args": "",
                "files": [{
                    "file_name": "meeting.docx",
                    "local_path": doc_path,
                }]
            }
        })
        bot_execute(msg, {})

        self.assertEqual(len(msg.outbox), 1)
        self.assertIn("【文档解析: meeting.docx】", msg.outbox[0].text)
        self.assertIn("会议纪要", msg.outbox[0].text)
        self.assertIn("讨论重点：提升系统鲁棒性。", msg.outbox[0].text)


if __name__ == "__main__":
    unittest.main()
