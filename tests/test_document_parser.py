import unittest
import os
import tempfile
from core.document_parser import parse_document, parse_pdf, parse_docx, parse_text_file


class TestDocumentParser(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()

    def test_parse_text_file(self):
        txt_path = os.path.join(self.temp_dir, "notes.txt")
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write("Line 1: 计划事项\nLine 2: 实施细节")

        res = parse_document(txt_path)
        self.assertTrue(res["ok"])
        self.assertEqual(res["type"], "text")
        self.assertIn("计划事项", res["content"])

    def test_parse_docx_file(self):
        import docx
        doc_path = os.path.join(self.temp_dir, "sample.docx")
        doc = docx.Document()
        doc.add_heading("项目方案", level=1)
        doc.add_paragraph("这是正文第一段。")
        
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "姓名"
        table.rows[0].cells[1].text = "职位"
        table.rows[1].cells[0].text = "张三"
        table.rows[1].cells[1].text = "架构师"
        
        doc.save(doc_path)

        res = parse_document(doc_path)
        self.assertTrue(res["ok"])
        self.assertEqual(res["type"], "docx")
        self.assertIn("# 项目方案", res["content"])
        self.assertIn("这是正文第一段。", res["content"])
        self.assertIn("| 姓名 | 职位 |", res["content"])
        self.assertIn("| 张三 | 架构师 |", res["content"])

    def test_parse_pdf_file(self):
        import pypdf
        pdf_path = os.path.join(self.temp_dir, "sample.pdf")
        
        # Create a simple PDF
        writer = pypdf.PdfWriter()
        page = writer.add_blank_page(width=200, height=200)
        
        with open(pdf_path, "wb") as f:
            writer.write(f)

        res = parse_document(pdf_path)
        self.assertTrue(res["ok"])
        self.assertEqual(res["type"], "pdf")
        self.assertEqual(res["page_count"], 1)


if __name__ == "__main__":
    unittest.main()
